# -*- coding: utf-8 -*-
"""
Vytvorí fiktívne testovacie dáta: MasterData_DEMO.xlsx a priečinok neroztriedene/.
Súbory sú už vygenerované – tento skript treba iba ak ich chceš vytvoriť znova.
Potrebuje: openpyxl, python-docx, reportlab, Pillow.
"""
import shutil
from pathlib import Path

from openpyxl import Workbook

TU = Path(__file__).resolve().parent
ZDROJ = TU / "neroztriedene"

# ---------------------------------------------------------------------------- MasterData
HARKY = {
    "Static": ["Tag No", "Part of scope", "Eq. Description", "Type"],
    "Rotating": ["Tag No", "Part of scope", "Eq. Description", "Type"],
    "Electric": ["Tag No", "Part of scope", "Eq. Description", "Type"],
    "Instrumentation": ["Tag No", "Part of scope", "Control Module", "I/O", "Assets No", "Eq. Description", "Type"],
    "Pipes": ["CL No.", "Technology ID", "Part of scope", "Assets No (SAP EQ)", "Pipeline name"],
}
DATA = {
    "Static": [
        ("11V109", "Y", "Odlučovač", "Nádrž vertikálna tlaková"),
        ("11V109A", "Y", "Odlučovač A", "Nádrž vertikálna tlaková"),
        ("11V109B", "Y", "Odlučovač B", "Nádrž vertikálna tlaková"),
        ("11V109C", "Y", "Odlučovač C", "Nádrž vertikálna tlaková"),
        ("11V510", "N", "Mimo rozsahu", "Nádrž vertikálna"),
        ("11E101", "Y", "Výmenník", "Výmenník rúrkový"),
        ("11E102", "Y", "Výmenník", "Výmenník rúrkový"),
        ("11A101", "Y", "Vzduchový chladič", "Výmenník vzduchový"),
        ("11A101A", "Y", "Vzduchový chladič A", "Výmenník vzduchový"),
        ("11A101B", "Y", "Vzduchový chladič B", "Výmenník vzduchový"),
        ("11A101B", "Y", "Vzduchový chladič B (duplicita v MasterData)", "Výmenník vzduchový"),
        ("11A101C", "Y", "Vzduchový chladič C", "Výmenník vzduchový"),
        ("11A101D", "Y", "Vzduchový chladič D", "Výmenník vzduchový"),
        ("11C101", "Y", "Kolóna", "Kolóna patrová"),
        ("11D101", "Y", "Zásobník", "Nádrž horizontálna"),
        ("11M101", "Y", "Mixér – teleso", "Mixér"),
        ("11F150A1", "Y", "Filter", "Filter kvapalín kovový"),
        ("E1", "Y", "Starý tag bez jednotky", "Výmenník"),
        ("E2", "Y", "Starý tag bez jednotky", "Výmenník"),
        ("PLUG-VALVE-1", "Y", "Ventil", "Ventil kužeľový"),
        ("PLUG-VALVE-12", "Y", "Ventil", "Ventil kužeľový"),
        ("3921_03", "Y", "Projekt 3921", "Iné"),
    ],
    "Rotating": [
        ("11P110A", "Y", "Čerpadlo A", "Čerpadlo"),
        ("11P110B", "Y", "Čerpadlo B", "Čerpadlo"),
        ("11M101", "Y", "Mixér – pohon", "Mixér"),
        ("P-11", "Y", "Starý tag", "Čerpadlo"),
    ],
    "Electric": [
        ("11E126A", "Y", "Elektrický ohrev", "Ohrev"),
        ("57DRM2-11E126A", "Y", "Vývod rozvádzača pre 11E126A", "Vývod rozvádzača"),
        ("ES11MA101A", "Y", "Elektro časť", "Elektromotor"),
        ("FAM11", "Y", "Rozvádzač", "Rozvádzač nn"),
    ],
    "Instrumentation": [
        ("11TE1347A1", "Y", None, None, None, "Snímač teploty", "Snímač teploty"),
        ("11TT1347A", "Y", None, None, None, "Vysielač", "Vysielač elektrický"),
        ("11PT1201", "Y", None, None, None, "Vysielač tlaku", "Snímač tlaku"),
        ("11XV1549", "N", None, None, None, "Ventil mimo rozsahu", "Ventil uzatvárací"),
    ],
    "Pipes": [
        (None, None, "Y", "PV900001", "11-DEA-N-1105-DN50-B1E (od 11P110A/B do"),
        (None, None, "Y", "PV900002", "11-AZF-230-DN100-B56A"),
        (None, None, "Y", "PV900003", "11-N2-0,6-1285-DN100-A1B"),
        (None, None, "Y", "PO900004", 'TBP-0279-3"-A1B'),
        (None, None, "Y", "PO900005", "Potrubný rozvod A/e-KSO"),
    ],
}


def masterdata():
    wb = Workbook()
    wb.remove(wb.active)
    for harok, hlavicka in HARKY.items():
        ws = wb.create_sheet(harok)
        ws.append([])                                        # ako originál: prázdny 1. riadok
        ws.append([None, "Názov objektu (SK hlavička)"])     # 2. riadok slovenská hlavička
        ws.append([None] + hlavicka)                         # 3. riadok anglická hlavička, od stĺpca B
        for r in DATA[harok]:
            ws.append([None] + list(r))
    wb.save(TU / "MasterData_DEMO.xlsx")


# ---------------------------------------------------------------------------- súbory
def pdf(cesta: Path, strany: list[str]):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    font = "Helvetica"
    for kandidat in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(kandidat).exists():
            pdfmetrics.registerFont(TTFont("Uni", kandidat))
            font = "Uni"
            break
    c = canvas.Canvas(str(cesta), pagesize=A4)
    for text in strany:
        y = 800
        c.setFont(font, 11)
        for riadok in text.splitlines():
            c.drawString(60, y, riadok)
            y -= 16
        c.showPage()
    c.save()


def sken_pdf(cesta: Path, text: str):
    """PDF iba s obrázkom (bez textovej vrstvy) – simuluje sken, ktorý bude potrebovať OCR."""
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    for i, riadok in enumerate(text.splitlines()):
        d.text((100, 120 + 40 * i), riadok, fill="black")
    img.save(cesta, "PDF", resolution=150)


def obrazok(cesta: Path, text: str):
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (800, 400), (235, 235, 225))
    ImageDraw.Draw(img).text((40, 180), text, fill="black")
    img.save(cesta, "JPEG")


def docx(cesta: Path, odseky: list[str], tabulka: list[list[str]] | None = None):
    from docx import Document
    d = Document()
    for o in odseky:
        d.add_paragraph(o)
    if tabulka:
        t = d.add_table(rows=len(tabulka), cols=len(tabulka[0]))
        for i, r in enumerate(tabulka):
            for j, v in enumerate(r):
                t.cell(i, j).text = v
    d.save(str(cesta))


def xlsx(cesta: Path, riadky: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Zoznam"
    for r in riadky:
        ws.append(r)
    wb.save(cesta)


def subory():
    if ZDROJ.exists():
        shutil.rmtree(ZDROJ)
    (ZDROJ / "revizie").mkdir(parents=True)
    (ZDROJ / "stare" / "kopia_z_mailu").mkdir(parents=True)
    (ZDROJ / "11D101").mkdir(parents=True)
    (ZDROJ / "rozne").mkdir(parents=True)

    # 1) revízna správa – pravidlo 3 (rozsah), 2 strany
    pdf(ZDROJ / "revizie" / "Revizna_sprava_2024_odlucovace.pdf", [
        "REVÍZNA SPRÁVA – VYHRADENÉ TLAKOVÉ ZARIADENIE\nTechn. Označenie: 11V109A\nNázov TZ: Odlučovač\n"
        "Rok výroby: 1978\nSkupina: A/b",
        "VN Vnútorná prehliadka nádob 11V109A-C vykonaná dňa 12. 3. 2024.\n"
        "Bez nálezu. Hrúbky stien v norme.\nST Skúška tesnosti – vyhovuje.",
    ])
    # 2) tá istá správa ešte raz v inom priečinku (duplicitný obsah -> 1 dokument)
    shutil.copy2(ZDROJ / "revizie" / "Revizna_sprava_2024_odlucovace.pdf",
                 ZDROJ / "stare" / "kopia_z_mailu" / "Revizna_sprava_2024_odlucovace (1).pdf")

    # 3) súpis zariadení – jeden dokument, veľa tagov (tu sa šetrí najviac kópií)
    xlsx(ZDROJ / "Supis_zariadeni_jednotka_11.xlsx", [
        ["Tag", "Popis", "Poznámka"],
        ["11V109", "Odlučovač", ""],
        ["11 V-109 A", "Odlučovač A", "zapísané s medzerami"],
        ["11V109B", "Odlučovač B", ""],
        ["11-E-101", "Výmenník", "zapísané s pomlčkami"],
        ["11E102", "Výmenník", ""],
        ["11A101A-D", "Vzduchové chladiče", "rozsah"],
        ["11C101", "Kolóna", ""],
        ["11D101", "Zásobník", ""],
        ["11M101", "Mixér", "teleso aj pohon"],
        ["11F150A1", "Filter", ""],
        ["11P110A/B", "Čerpadlá", "lomka = A a B"],
        ["11TE1347A1", "Snímač", ""],
        ["11TT1347A", "Vysielač", ""],
        ["11PT-1201", "Vysielač tlaku", ""],
        ["11-AZF-230-DN100-B56A", "Potrubie", ""],
    ])

    # 4) protokol vo Worde – spojka 'a', rozsah, potrubie s 'DN 50'
    docx(ZDROJ / "revizie" / "Protokol_tlakova_skuska_2023.docx", [
        "PROTOKOL O TLAKOVEJ SKÚŠKE",
        "Skúšané zariadenie 11A101 a pripojené potrubie 11-DEA-N-1105-DN 50-B1E.",
        "Výmenníky 11E101 a 11E102 neboli súčasťou skúšky.",
        "Skúšobný tlak 1,5 MPa, výsledok vyhovuje.",
    ], [["Zariadenie", "Výsledok"], ["11A101", "vyhovuje"]])

    # 5) krátky tag v názve súboru (automaticky) a v obsahu (na kontrolu)
    (ZDROJ / "rozne" / "E1_kalibracia_2022.txt").write_text(
        "Kalibrácia snímača. Porovnanie s E2 podľa postupu QP-07.\n", encoding="utf-8")

    # 6) poznámky – tag bez predpony (na kontrolu), PLUG-VALVE (nejednoznačný), hranice
    (ZDROJ / "rozne" / "poznamky_z_obchodzky.txt").write_text(
        "Obchôdzka 4. 5. 2024\n"
        "- na V109A presakuje príruba (bez predpony jednotky)\n"
        "- PLUG-VALVE-1 ťažko ide\n"
        "- 11V1090 a 11E1011 nie sú naše zariadenia (test hraníc)\n"
        "- nádrž 11V109 v prevádzke bez problémov\n", encoding="utf-8")

    # 7) elektro – tag vnútri dlhšieho tagu (pravidlo 5)
    (ZDROJ / "rozne" / "Elektro_vyvody_rozvadzac_DRM2.csv").write_text(
        "Vyvod;Popis\n57DRM2-11E126A;Vývod rozvádzača pre ohrev\nES11MA101A;Elektro časť motora\n",
        encoding="cp1250")

    # 8) prázdny súbor, tag iba v názve (simulácia 'nečitateľného' súboru)
    (ZDROJ / "rozne" / "11P110A_meranie_vibracii.pdf").write_bytes(b"")

    # 9) sken bez textu -> potrebuje OCR, tag nájdený iba z názvu
    sken_pdf(ZDROJ / "revizie" / "sken_protokol_11C101.pdf", "PROTOKOL\n11C101\nvnutorna prehliadka")

    # 10) fotka štítku -> potrebuje OCR, bez tagu v názve
    obrazok(ZDROJ / "rozne" / "foto_stitok_IMG_2041.jpg", "11E102  vyrobny stitok")

    # 11) potrubie – palce vs. DN, potrubie s 0,6 a hromadný rozvod
    (ZDROJ / "rozne" / "potrubia_TBP_merania_hrubok.txt").write_text(
        "Meranie hrúbok:\nTBP-0279-DN80-A1B  min 4,1 mm\n11-N2-0,6-1285-DN100-A1B  min 3,8 mm\n"
        "Potrubný rozvod A/e-KSO – vizuálna kontrola OK\n", encoding="utf-8")

    # 12) zariadenie mimo rozsahu (Part of scope = N) -> nenájde sa
    (ZDROJ / "rozne" / "zariadenie_mimo_rozsahu.txt").write_text(
        "Oprava nádrže 11V510 a ventilu 11XV1549.\n", encoding="utf-8")

    # 13) tag v názve priečinka
    (ZDROJ / "11D101" / "fotodokumentacia_popis.txt").write_text(
        "Fotky z odstávky 2023, pohľad zboku a zhora.\n", encoding="utf-8")

    # 14) RTF (starý formát)
    (ZDROJ / "revizie" / "stary_zapis_1998.rtf").write_text(
        r"{\rtf1\ansi\ansicpg1250 {\fonttbl{\f0 Arial;}}\f0 Z\'e1pis o oprave zariadenia 11D101 a 11F150A1.\par}",
        encoding="latin-1")

    # 15) dočasný súbor Wordu (ignoruje sa) a nepodporovaný typ
    (ZDROJ / "revizie" / "~$otokol_tlakova_skuska_2023.docx").write_bytes(b"\x00docasny")
    (ZDROJ / "rozne" / "vykres_11E101.dwg").write_bytes(b"AC1027 fiktivny vykres")


# ---------------------------------------------------------------------------- realistické skeny (OCR)
def _font(velkost: int):
    from PIL import ImageFont
    for f in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"):
        if Path(f).exists():
            return ImageFont.truetype(f, velkost)
    return ImageFont.load_default()


def _sken_strana(riadky: list[str], uhol: float = 0.0, otocit_90: bool = False, dpi: int = 250):
    """A4 strana ako zo skenera: písmo ~11 pt, mierne natočenie, šum, sivé pozadie, jemné rozmazanie."""
    import random
    from PIL import Image, ImageDraw, ImageFilter
    random.seed(len("".join(riadky)))
    w, h = int(8.27 * dpi), int(11.69 * dpi)
    img = Image.new("L", (w, h), 248)
    d = ImageDraw.Draw(img)
    font = _font(int(dpi * 11 / 72))
    nadpis = _font(int(dpi * 15 / 72))
    y = int(1.2 * dpi)
    for i, r in enumerate(riadky):
        d.text((int(1.0 * dpi), y), r, fill=25, font=nadpis if i == 0 else font)
        y += int(dpi * (0.42 if i == 0 else 0.30))
    if otocit_90:
        img = img.rotate(90, expand=True, fillcolor=248)
    img = img.rotate(uhol, resample=Image.BICUBIC, expand=False, fillcolor=248)
    px = img.load()
    for _ in range(w * h // 900):                      # šum (prach zo skenera)
        x, yy = random.randrange(img.width), random.randrange(img.height)
        px[x, yy] = random.choice((120, 180, 255))
    return img.filter(ImageFilter.GaussianBlur(0.6))


def _pdf_zo_stran(cesta: Path, strany: list):
    """Strany = PIL obrázok (sken) alebo str (digitálny text). Výsledok: PDF bez textovej vrstvy pri skenoch."""
    import io
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    font = "Helvetica"
    for kandidat in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(kandidat).exists():
            pdfmetrics.registerFont(TTFont("Uni", kandidat))
            font = "Uni"
            break
    c = canvas.Canvas(str(cesta), pagesize=A4, invariant=1)
    for s in strany:
        if isinstance(s, str):
            c.setPageSize(A4)
            y = 790
            for i, r in enumerate(s.splitlines()):
                c.setFont(font, 15 if i == 0 else 11)
                c.drawString(60, y, r)
                y -= 22 if i == 0 else 16
        else:
            velkost = landscape(A4) if s.width > s.height else A4
            c.setPageSize(velkost)
            buf = io.BytesIO()
            s.save(buf, "JPEG", quality=55)
            buf.seek(0)
            c.drawImage(ImageReader(buf), 0, 0, width=velkost[0], height=velkost[1])
        c.showPage()
    c.save()


def skeny():
    """Realistické skeny na test OCR (fiktívny obsah). Nemení ostatné demo súbory."""
    from PIL import Image, ImageDraw, ImageFilter
    ciel = ZDROJ / "skeny"
    ciel.mkdir(parents=True, exist_ok=True)

    # A) 2-stranový sken revíznej správy – tag, rozsah, ďalší tag; mierne natočený
    _pdf_zo_stran(ciel / "Revizia_vymennikov_sken_2019.pdf", [
        _sken_strana(["REVÍZNA SPRÁVA č. 47/2019 (FIKTÍVNA – DEMO)",
                      "Prevádzkovateľ: Demo rafinéria, a. s.",
                      "Techn. označenie: 11E101",
                      "Názov TZ: Výmenník tepla rúrkový",
                      "Rok výroby: 1976        Skupina: A/b",
                      "Výsledok revízie: vyhovuje"], uhol=1.2),
        _sken_strana(["Pokračovanie – strana 2",
                      "Vnútorná prehliadka vzduchových chladičov 11A101A-C bez nálezu.",
                      "Tlaková skúška výmenníka 11E102 – vyhovuje.",
                      "Hrúbky stien merané ultrazvukom, minimum 6,2 mm."], uhol=-0.8),
    ])

    # B) zmiešané PDF: strana 1 digitálny text, strana 2 sken (čiastočné OCR)
    #    zámerná chyba v zápise 11VI09B (písmeno I namiesto 1) -> ukážka „OCR oprava“ na kontrolu
    _pdf_zo_stran(ciel / "Zmiesany_protokol_oprava_2021.pdf", [
        "PROTOKOL O OPRAVE (FIKTÍVNY – DEMO)\nTextová strana vytvorená v PC.\n"
        "Opravované zariadenie: kolóna 11C101\nDátum opravy: 14. 6. 2021",
        _sken_strana(["Príloha – naskenovaný formulár",
                      "Vymenené tesnenie príruby na nádobe 11VI09B",
                      "(v pôvodnom zápise je písmeno I namiesto jednotky)",
                      "Kontrola zásobníka 11D101 bez nálezu."], uhol=0.6),
    ])

    # C) sken otočený o 90° (OCRmyPDF ho narovná – --rotate-pages)
    _pdf_zo_stran(ciel / "Kalibracia_otoceny_sken_2018.pdf", [
        _sken_strana(["KALIBRAČNÝ LIST (FIKTÍVNY – DEMO)",
                      "Filter 11F150A1 – snímač tlakovej diferencie",
                      "Dátum kalibrácie: 3. 4. 2018",
                      "Výsledok: v tolerancii"], otocit_90=True),
    ])

    # D) fotka výrobného štítku (obrázok, bez tagu v názve)
    img = Image.new("RGB", (1800, 1100), (128, 122, 110))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((180, 160, 1620, 940), radius=30, fill=(214, 214, 206), outline=(90, 90, 90), width=6)
    for (x, y), t, f in (((260, 230), "VÝROBNÝ ŠTÍTOK", _font(80)), ((260, 400), "Čerpadlo 11P110B", _font(96)),
                         ((260, 560), "Q = 25 m3/h    H = 48 m", _font(64)),
                         ((260, 680), "Rok výroby: 1981  (DEMO)", _font(64))):
        d.text((x, y), t, fill=(30, 30, 30), font=f)
    img = img.rotate(-3, resample=Image.BICUBIC, fillcolor=(128, 122, 110)).filter(ImageFilter.GaussianBlur(1.0))
    img.save(ciel / "foto_stitok_cerpadla.jpg", "JPEG", quality=80)


if __name__ == "__main__":
    import sys
    if "--iba-skeny" not in sys.argv:
        masterdata()
        subory()
    skeny()
    print("Demo dáta vytvorené v", TU)
