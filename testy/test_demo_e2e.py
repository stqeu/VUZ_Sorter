# -*- coding: utf-8 -*-
"""End-to-end test celého behu na fiktívnych demo dátach.

Overuje: TEST režim nič nevytvorí, ostrý beh uloží každý dokument raz a vytvorí odkazy,
opakovaný beh nič nezduplikuje.  Spustenie:  python -m unittest discover testy
"""

import configparser
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

KOREN = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KOREN))
import lta_triedenie as lt  # noqa: E402


def _suhrn(xlsx: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True)
    try:
        return {r[0]: r[1] for r in wb["Súhrn"].iter_rows(min_row=2, values_only=True)}
    finally:
        wb.close()


class TestDemoBeh(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="vuz_sorter_"))
        shutil.copytree(KOREN / "demo_data", cls.tmp / "demo_data")
        cp = configparser.ConfigParser(inline_comment_prefixes=(";", "#"))
        cp.read(KOREN / "config" / "demo.ini", encoding="utf-8")
        cp["projekt"]["masterdata"] = str(cls.tmp / "demo_data" / "MasterData_DEMO.xlsx")
        cp["projekt"]["zdroje"] = str(cls.tmp / "demo_data" / "neroztriedene")
        cp["projekt"]["vystup"] = str(cls.tmp / "vystup")
        cp["projekt"]["typ_odkazu"] = "url"          # funguje na každom OS bez PowerShellu
        cls.ini = cls.tmp / "test.ini"
        with cls.ini.open("w", encoding="utf-8") as f:
            cp.write(f)
        cls.vystup = cls.tmp / "vystup"

    @classmethod
    def tearDownClass(cls):
        import logging
        for h in list(logging.getLogger().handlers):    # uvoľniť log súbor (Windows)
            h.close()
            logging.getLogger().removeHandler(h)
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_1_test_rezim_nic_nevytvori(self):
        self.assertEqual(lt.main(["--config", str(self.ini)]), 0)
        self.assertTrue((self.vystup / "protokol_TEST.xlsx").is_file())
        self.assertFalse((self.vystup / "_dokumenty").exists())
        self.assertFalse((self.vystup / "zariadenia").exists())
        s = _suhrn(self.vystup / "protokol_TEST.xlsx")
        self.assertEqual(s["Unikátnych dokumentov (podľa obsahu)"], 14)
        self.assertEqual(s["Duplicitné kópie v zdrojoch"], 1)
        self.assertEqual(s["Stav: auto"], 39)
        self.assertEqual(s["Stav: na kontrolu"], 4)

    def test_2_ostry_beh_a_opakovanie(self):
        self.assertEqual(lt.main(["--config", str(self.ini), "--ostro"]), 0)
        s = _suhrn(self.vystup / "index.xlsx")
        self.assertEqual(s["Nové kópie do úložiska"], 12)
        self.assertEqual(s["Vytvorené / aktualizované odkazy"], 40)
        self.assertEqual(len(list((self.vystup / "_dokumenty").iterdir())), 12)
        odkazy = list((self.vystup / "zariadenia").rglob("*.url"))
        self.assertEqual(len(odkazy), 40)
        self.assertEqual({o.name for o in (self.vystup / "zariadenia" / "Static" / "11V109A").iterdir()},
                         {"Revizna_sprava_2024_odlucovace.pdf.url", "Supis_zariadeni_jednotka_11.xlsx.url"})

        # opakovaný beh – nič nové
        self.assertEqual(lt.main(["--config", str(self.ini), "--ostro"]), 0)
        s = _suhrn(self.vystup / "index.xlsx")
        self.assertEqual(s["Nové kópie do úložiska"], 0)
        self.assertEqual(s["Vytvorené / aktualizované odkazy"], 0)
        self.assertEqual(s["Zmazané zastarané odkazy"], 0)


if __name__ == "__main__":
    unittest.main()
